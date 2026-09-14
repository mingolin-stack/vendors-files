import os
import io
import fitz  # PyMuPDF
import docx
import pandas as pd
import json
import base64
from openai import OpenAI
import streamlit as st

def extract_content_from_file(file_bytes: bytes, file_ext: str):
    """根據副檔名提文字或轉成圖片 bytes"""
    file_ext = file_ext.lower()
    
    if file_ext == ".pdf":
        # 將 PDF 第一頁轉為 PNG 圖片
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        if len(doc) == 0:
            raise ValueError("上傳的 PDF 檔案為空檔")
        page = doc[0]
        pix = page.get_pixmap(dpi=200)
        return "image", pix.tobytes("png")

    elif file_ext in [".docx", ".doc"]:
        # 提取 Word 文件中的文字與表格
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_data:
                    full_text.append(" | ".join(row_data))
        return "text", "
".join(full_text)

    elif file_ext in [".xlsx", ".xls"]:
        # 讀取 Excel 並轉為 CSV 文字字串
        df_dict = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None)
        all_sheets_text = []
        for sheet_name, df in df_dict.items():
            all_sheets_text.append(f"--- Sheet: {sheet_name} ---")
            all_sheets_text.append(df.to_csv(index=False))
        return "text", "
".join(all_sheets_text)

    else:
        raise ValueError(f"不支援的檔案格式: {file_ext}")


def analyze_vendor_document(file_bytes: bytes, file_ext: str) -> dict:
    """分析 Word, Excel, PDF 文件並提取供應商欄位資訊"""
    api_key = st.secrets.get("OPENAI_API_KEY")
    if not api_key:
        raise ValueError("未設定 OPENAI_API_KEY，請在 Streamlit Secrets 設定中新增。")

    client = OpenAI(api_key=api_key)
    content_type, extracted_content = extract_content_from_file(file_bytes, file_ext)

    prompt = """
    請幫我辨識這份供應商資料文件的內容，並以 JSON 格式回傳以下欄位：
    - 公司全名
    - 統一編號
    - 負責人
    - 聯絡人
    - 公司電話
    - 聯絡電話
    - 聯絡地址
    - 帳單地址
    - 匯款帳號戶名
    - 匯款銀行
    - 分行別
    - 匯款帳號

    若欄位不存在或無法辨識，請填寫 ""。僅回傳純 JSON 字串，不要包含任何 markdown 標籤。
    """

    if content_type == "image":
        base64_image = base64.b64encode(extracted_content).decode('utf-8')
        response = client.chat.completions.create(
            model="gpt-4o",
            messages=[
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": prompt},
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:image/png;base64,{base64_image}"
                            },
                        },
                    ],
                }
            ],
            max_tokens=1000,
        )
    else:
        user_message = f"{prompt}

以下是文件的內容資料：
{extracted_content}"
        response = client.chat.completions.create(
            model="gpt-4o",
            messages=[
                {"role": "user", "content": user_message}
            ],
            max_tokens=1000,
        )

    result_text = response.choices[0].message.content.strip()
    if result_text.startswith("```json"):
        result_text = result_text[7:]
    if result_text.endswith("```"):
        result_text = result_text[:-3]

    return json.loads(result_text.strip())
